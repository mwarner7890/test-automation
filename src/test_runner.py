import adb as adb_module
import csv
import os
import standard_testing
import subprocess
import throughput_testing
import shutil
import sys
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from throughput_ftp import ThroughputFTP


def set_up(adb_instance):
    adb_instance.start_server()
    adb_instance.unlock_screen_with_pin()


def tear_down(adb_instance):
    adb_instance.input_key_event_sequence(['KEYCODE_HOME',
                                          'KEYCODE_POWER'])


def run_tests(**kwargs):
    adb_instances = [kwargs.get('adb_device_1'),
                     kwargs.get('adb_device_2')]
    num_of_tests = len(test_names)
    test_count = 1
    passed_count = 0
    failed_count = 0
    running_throughput_test = sys.argv[1] == 'throughput_test'
    results_dir = os.path.join(os.path.expanduser('~'), 'Documents', 'throughput_test',
                               'results')
    for test_name in test_names:
        print('Running test ({}/{}): {}'.format(test_count,
                                                num_of_tests,
                                                test_name))
        try:
            if running_throughput_test:
                test_method = getattr(throughput_testing, test_name)
            else:
                test_method = getattr(standard_testing, test_name)
        except AttributeError:
            print('Error: No test case named {}'.format(test_name))
            for adb_instance in adb_instances:
                if adb_instance:
                    tear_down(adb_instance)
            num_of_tests -= 1
            continue
        if running_throughput_test:
            csv_filename = test_names[0] + '.csv'
            xlsx_filename = 'throughput_results.xlsx'
            _create_csv_file_for_results(results_dir, csv_filename, adb_device_1.model_name,
                                         adb_device_2.model_name)
            _copy_results_xlsx_template_if_not_exists(results_dir, xlsx_filename)
            xlsx_results = [adb_instances[0].model_name, adb_instances[1].model_name]
            ftp_config_dir = os.path.expanduser('~') + '/Documents/throughput_test'
            ftp_config_fname = 'ftp_config.json'
            ftp = ThroughputFTP(ftp_config_dir, ftp_config_fname)
            for _ in range(0, ftp.test_count):
                current_csv_result_line = []
                for adb_instance in adb_instances:
                    if adb_instance:
                        set_up(adb_instance)
                        adb_instance.ftp = ftp
                        temp_result = test_method(adb_instance)
                        if temp_result == 'Fail':
                            failed_count += 1
                        else:
                            passed_count += 1
                        current_csv_result_line.append(temp_result)
                        xlsx_results.append(temp_result)
                        tear_down(adb_instance)
                _add_line_to_results_csv(results_dir, current_csv_result_line, csv_filename)

            _save_results_to_xlsx(results_dir, xlsx_filename, xlsx_results,
                                  test_name.split('_')[1], adb_instances)
        else:
            set_up(adb_instances[0])
            test_passed = test_method(adb_instances[0])

            if test_passed:
                passed_count += 1
            else:
                failed_count += 1
            print(eval_test_result(test_passed))

        for adb_instance in adb_instances:
            if adb_instance and not running_throughput_test:
                tear_down(adb_instance)
        test_count += 1

    if running_throughput_test:
        print('Finished running {} tests(s):'.format(num_of_tests))
        print('{} passed'.format(passed_count))
        print('{} failed'.format(failed_count))


def eval_test_result(result):
    if result:
        return 'Test passed'
    return 'Test failed'


def _create_csv_file_for_results(results_dir, csv_filename, device_model_name_1,
                                 device_model_name_2):
    _create_results_dir_if_not_exists(results_dir)

    with open(os.path.join(results_dir, csv_filename), 'w+',
              newline="\n", encoding="utf-8") as csvfile:
        csvwriter = csv.writer(csvfile, delimiter=',')
        csvwriter.writerow([device_model_name_1, device_model_name_2])


def _create_results_dir_if_not_exists(results_dir):
    if not os.path.isdir(results_dir):
        os.makedirs(results_dir)


def _add_line_to_results_csv(results_dir, csv_line, csv_filename):
    with open(os.path.join(results_dir, csv_filename), 'a',
              newline="\n", encoding="utf-8") as csvfile:
        csvwriter = csv.writer(csvfile, delimiter=',')
        csvwriter.writerow([csv_line[0], csv_line[1]])


def _copy_results_xlsx_template_if_not_exists(results_dir, xlsx_filename):
    xlsx_fpath = os.path.join(results_dir, xlsx_filename)
    if not os.path.isfile(xlsx_fpath):
        shutil.copy('../templates/throughput_results_template.xlsx',
                    xlsx_fpath)


def _coord_to_x_y(coord):
    coord = str.upper(coord)
    return ord(coord[0]) - 64, int(coord[1:])


def _create_result_chart_at_position(position, results_range, chart_title, workbook):
    worksheet = workbook.active

    chart = BarChart()
    chart.title = chart_title
    chart.y_axis.title = 'Download time (seconds)'

    min_col, min_row = _coord_to_x_y(results_range[0])
    max_col, max_row = _coord_to_x_y(results_range[1])

    values = Reference(worksheet, min_col=min_col, min_row=min_row,
                       max_col=max_col, max_row=max_row)
    chart.add_data(values, titles_from_data=True)

    worksheet.add_chart(chart, position)


def _save_results_to_xlsx(results_dir, xlsx_filename, results, test_type,
                          adb_instances):
    _create_results_dir_if_not_exists(results_dir)
    xlsx_fpath = os.path.join(results_dir, xlsx_filename)

    spreadsheet_cell_addresses = {
        '2g':
            {
                'results_cell_range': ('B2', 'C12'),
                'result_chart_position': 'T2',
                'device_info_cells':
                [
                    {
                        'sim_operator_cell': 'B15',
                        'build_cell': 'B16',
                        'sw_version_cell': 'B17',
                        'android_version_cell': 'B18',
                        'manufacturer_cell': 'B19',
                        'imei_cell': 'B20',
                        'result_chart_position': 'K2'
                    },
                    {
                        'sim_operator_cell': 'C15',
                        'build_cell':  'C16',
                        'sw_version_cell': 'C17',
                        'android_version_cell': 'C18',
                        'manufacturer_cell': 'C19',
                        'imei_cell': 'C20'
                    }
                ]
            },
        '3g':
            {
                'results_cell_range': ('E2', 'F12'),
                'result_chart_position': 'T21',
                'device_info_cells':
                [
                    {
                        'sim_operator_cell': 'E15',
                        'build_cell':  'E16',
                        'sw_version_cell': 'E17',
                        'android_version_cell': 'E18',
                        'manufacturer_cell': 'E19',
                        'imei_cell': 'E20',
                        'result_chart_position': 'K21'
                    },
                    {
                        'sim_operator_cell': 'F14',
                        'build_cell': 'F15',
                        'sw_version_cell': 'F17',
                        'android_version_cell': 'F18',
                        'manufacturer_cell': 'F19',
                        'imei_cell': 'F20'
                    }
                ]
            },
        '4g':
            {
                'results_cell_range': ('H2', 'I12'),
                'result_chart_position': 'T33',
                'device_info_cells':
                [
                    {
                        'sim_operator_cell': 'H15',
                        'build_cell': 'H16',
                        'sw_version_cell': 'H17',
                        'android_version_cell': 'H18',
                        'manufacturer_cell': 'H19',
                        'imei_cell': 'H20',
                        'result_chart_position': 'K33'
                    },
                    {
                        'sim_operator_cell': 'I15',
                        'build_cell': 'I16',
                        'sw_version_cell': 'I17',
                        'android_version_cell': 'I18',
                        'manufacturer_cell': 'I19',
                        'imei_cell': 'I20'
                    }
                ]
            },
        'wifi':
            {
                'results_cell_range': ('B22', 'C32'),
                'result_chart_position': 'K52',
                'device_info_cells':
                [
                    {
                        'sim_operator_cell': 'B35',
                        'build_cell': 'B36',
                        'sw_version_cell': 'B37',
                        'android_version_cell': 'B38',
                        'manufacturer_cell': 'B39',
                        'imei_cell': 'B40'
                    },
                    {
                        'sim_operator_cell': 'C35',
                        'build_cell': 'C36',
                        'sw_version_cell': 'C37',
                        'android_version_cell': 'C38',
                        'manufacturer_cell': 'C39',
                        'imei_cell': 'C40'
                    }
                ]
            }
    }

    results_workbook = load_workbook(xlsx_fpath)
    results_worksheet = results_workbook.active

    cells = []
    for cell_tuple in \
            results_worksheet[spreadsheet_cell_addresses[test_type]['results_cell_range'][0]:spreadsheet_cell_addresses[test_type]['results_cell_range'][1]]:
        for cell in cell_tuple:
            cells.append(cell)

    for cell, result in zip(cells, results):
        cell.value = result

    for adb_instance, cell_addresses in zip(adb_instances, spreadsheet_cell_addresses[test_type]['device_info_cells']):
        results_worksheet[cell_addresses['sim_operator_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.prop.get('gsm.sim.operator.alpha'))
        results_worksheet[cell_addresses['build_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.prop['ro.build.flavor'])
        results_worksheet[cell_addresses['sw_version_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.prop.get('ro.build.version.software'))
        results_worksheet[cell_addresses['android_version_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.prop.get('ro.build.version.release'))
        results_worksheet[cell_addresses['manufacturer_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.prop.get('ro.product.manufacturer'))
        results_worksheet[cell_addresses['imei_cell']].value = \
            _strip_r_n_and_remove_comma(adb_instance.get_imei())

    _create_result_chart_at_position(spreadsheet_cell_addresses[test_type]['result_chart_position'],
                                     spreadsheet_cell_addresses[test_type]['results_cell_range'],
                                     str.upper(test_type) + ' Test Results', results_workbook)
    results_workbook.save(xlsx_fpath)


def _parse_cmd_line_args(args):  # TODO: simplify this!!!
    parsed_test_names = []
    parse_error = ''
    usage_msg = 'Usage:\n' \
                'To run standard tests (single device): standard_tests\n' \
                'To run specific standard test(s): test_only <test1> <test2>\n' \
                'To run throughput tests (between two devices):\n' \
                '    All throughput testing: throughput_test all\n' \
                '    Specific RAT/Wi-Fi testing: throughput_test <2g/3g/4g/wifi>'

    if len(args) == 1:
        parse_error = usage_msg
    else:
        if args[1] == 'test_only':
            if args[2:]:
                parsed_test_names = args[2:]
            else:
                parse_error = 'Please specify the tests to run.\n' \
                              'Usage: test_only test1 test2'
        elif args[1] == 'standard_tests':
            parsed_test_names = [obj for obj in dir(standard_testing) if 'test_' in obj]
        elif args[1] == 'throughput_test':
            if len(args) < 3:
                parse_error = 'Please specify the RAT or WiFi.\n'\
                    'Usage: throughput_test 2g/3g/4g/wifi'
            else:
                if args[2] == 'all':
                    parsed_test_names = [obj for obj in dir(throughput_testing) if 'test_' in obj]
                else:
                    if len(args) > 3:
                        for parsed_test_type in args[2:]:
                            parsed_test_names.append('test_{}_throughput'.format(parsed_test_type))
                    else:
                        if args[2] not in ['2g', '3g', '4g', 'wifi']:
                            parse_error = 'Please specify the RAT or WiFi.\n'\
                                          'Usage: throughput_test 2g/3g/4g/wifi'
                        else:
                            parsed_test_names = ['test_{}_throughput'.format(args[2])]
        else:
            parse_error = 'Unknown option "{}"\n' \
                            '{}'.format(args[1], usage_msg)

    return parsed_test_names, parse_error


def _get_device_names():
    raw_output = adb_module.get_stdout_from_command('adb devices')
    return raw_output.replace('\\tdevice', '').split('\\r\\n')[1:][:-2]


def _get_device_model_name(device_name):
    model_name = adb_module.get_stdout_from_command(
        'adb -s {} shell getprop ro.product.model'.format(device_name))
    return model_name.replace('b\'', '').replace('\\r', '').replace('\\n', '')[:-1]


def _strip_r_n_and_remove_comma(string):
    if string:
        s = string.replace('\\r', '').replace('\\n', '')
        if s == '[,]' or s == ',':
            return ''
        return s
    return ''


if __name__ == '__main__':
    test_names, error = _parse_cmd_line_args(sys.argv)
    if error:
        print(error)
    else:
        if 'throughput' in test_names[0]:
            try:
                device_name_1, device_name_2 = _get_device_names()
                adb_device_1 = adb_module.Adb(device_name=device_name_1,
                                              model_name=_get_device_model_name(device_name_1))
                adb_device_2 = adb_module.Adb(device_name=device_name_2,
                                              model_name=_get_device_model_name(device_name_2))
                run_tests(adb_device_1=adb_device_1, adb_device_2=adb_device_2)
            except ValueError:
                print('Please connect two devices for throughput testing',
                      file=sys.stderr)
            except KeyboardInterrupt:
                print('Stopping setup... (interrupted by user)')

        else:
            try:
                run_tests(adb_device_1=adb_module.Adb())
            except subprocess.CalledProcessError:
                print('Please connect one device for standard tests',
                      file=sys.stderr)
            except KeyboardInterrupt:
                print('Stopping tests... (interrupted by user)')
